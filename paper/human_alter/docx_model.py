from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SYNC_PREFIX = "SYNC_"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def normalized_text(value: str, *, compact_cjk: bool = False, casefold: bool = False) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\u00ad", "").replace("\u2011", "-").replace("\u2013", "-").replace("\u2014", "-")
    value = re.sub(r"\s+", " ", value).strip()
    if compact_cjk:
        value = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fffA-Za-z0-9])", "", value)
        value = re.sub(r"(?<=[A-Za-z0-9])\s+(?=[\u3400-\u9fff])", "", value)
    return value.casefold() if casefold else value


def text_hash(value: str, **kwargs) -> str:
    return sha256_bytes(normalized_text(value, **kwargs).encode("utf-8"))


def paragraph_text(paragraph) -> str:
    out = []
    for node in paragraph._p.iter():
        if node.tag == W + "t":
            out.append(node.text or "")
        elif node.tag == W + "tab":
            out.append("\t")
        elif node.tag in {W + "br", W + "cr"}:
            out.append("\n")
        elif node.tag == W + "noBreakHyphen":
            out.append("-")
    return "".join(out)


def bookmark_names(paragraph) -> list[str]:
    return [node.get(qn("w:name")) for node in paragraph._p.iter(W + "bookmarkStart")]


def next_bookmark_id(doc) -> int:
    values = []
    for node in doc.element.body.iter(W + "bookmarkStart"):
        try:
            values.append(int(node.get(qn("w:id"))))
        except (TypeError, ValueError):
            pass
    return max(values, default=0) + 1


def anchor_name(prefix: str, identifier: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", identifier)
    raw = f"SYNC_{prefix}_{safe}"
    if len(raw) <= 40:
        return raw
    digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:8]
    return f"SYNC_{prefix}_{safe[:26]}_{digest}"[:40]


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    if name in bookmark_names(paragraph):
        return
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    children = list(paragraph._p)
    pos = 1 if children and children[0].tag == W + "pPr" else 0
    paragraph._p.insert(pos, start)
    paragraph._p.append(end)


def find_anchor_paragraph(doc, name: str):
    matches = [p for p in doc.paragraphs if name in bookmark_names(p)]
    if len(matches) != 1:
        return None
    return matches[0]


def paragraph_index(doc, paragraph) -> int:
    target = paragraph._p
    for index, candidate in enumerate(doc.paragraphs):
        if candidate._p is target:
            return index
    return -1


def _sync_bookmark_nodes(paragraph):
    ids = {
        node.get(qn("w:id"))
        for node in paragraph._p.iter(W + "bookmarkStart")
        if (node.get(qn("w:name")) or "").startswith(SYNC_PREFIX)
    }
    return [
        node for node in paragraph._p
        if (node.tag == W + "bookmarkStart" and node.get(qn("w:id")) in ids)
        or (node.tag == W + "bookmarkEnd" and node.get(qn("w:id")) in ids)
    ]


def copy_paragraph_content(source, target) -> None:
    anchors = [deepcopy(node) for node in _sync_bookmark_nodes(target)]
    for child in list(target._p):
        if child.tag != W + "pPr":
            target._p.remove(child)
    starts = [node for node in anchors if node.tag == W + "bookmarkStart"]
    ends = [node for node in anchors if node.tag == W + "bookmarkEnd"]
    for node in starts:
        target._p.append(node)
    source_sync_ids = {
        node.get(qn("w:id"))
        for node in source._p.iter(W + "bookmarkStart")
        if (node.get(qn("w:name")) or "").startswith(SYNC_PREFIX)
    }
    for child in source._p:
        if child.tag == W + "pPr":
            continue
        if child.tag in {W + "bookmarkStart", W + "bookmarkEnd"} and child.get(qn("w:id")) in source_sync_ids:
            continue
        target._p.append(deepcopy(child))
    for node in ends:
        target._p.append(node)


def set_paragraph_text(paragraph, value: str) -> None:
    anchors = [deepcopy(node) for node in _sync_bookmark_nodes(paragraph)]
    first_rpr = None
    for run in paragraph._p.iter(W + "r"):
        rpr = run.find(W + "rPr")
        if rpr is not None:
            first_rpr = deepcopy(rpr)
            break
    for child in list(paragraph._p):
        if child.tag != W + "pPr":
            paragraph._p.remove(child)
    for node in [x for x in anchors if x.tag == W + "bookmarkStart"]:
        paragraph._p.append(node)
    run = OxmlElement("w:r")
    if first_rpr is not None:
        run.append(first_rpr)
    text = OxmlElement("w:t")
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = value
    run.append(text)
    paragraph._p.append(run)
    for node in [x for x in anchors if x.tag == W + "bookmarkEnd"]:
        paragraph._p.append(node)


def copy_paragraph_group_content(sources, target) -> None:
    """Copy multiple source paragraphs into one target with line breaks."""
    anchors = [deepcopy(node) for node in _sync_bookmark_nodes(target)]
    for child in list(target._p):
        if child.tag != W + "pPr":
            target._p.remove(child)
    for node in [x for x in anchors if x.tag == W + "bookmarkStart"]:
        target._p.append(node)
    for index, source in enumerate(sources):
        if index:
            run = OxmlElement("w:r")
            run.append(OxmlElement("w:br"))
            target._p.append(run)
        source_sync_ids = {
            node.get(qn("w:id"))
            for node in source._p.iter(W + "bookmarkStart")
            if (node.get(qn("w:name")) or "").startswith(SYNC_PREFIX)
        }
        for child in source._p:
            if child.tag == W + "pPr":
                continue
            if child.tag in {W + "bookmarkStart", W + "bookmarkEnd"} and child.get(qn("w:id")) in source_sync_ids:
                continue
            target._p.append(deepcopy(child))
    for node in [x for x in anchors if x.tag == W + "bookmarkEnd"]:
        target._p.append(node)


def table_matrix(table) -> list[list[str]]:
    return [[normalized_text(cell.text) for cell in row.cells] for row in table.rows]


def table_hash(table) -> str:
    return sha256_bytes(json.dumps(table_matrix(table), ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def replace_table(source, target) -> None:
    target._tbl.getparent().replace(target._tbl, deepcopy(source._tbl))


def image_parts(doc) -> list[tuple[str, object]]:
    output = []
    for blip in doc.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        rid = blip.get(R + "embed")
        if rid and rid in doc.part.related_parts:
            output.append((rid, doc.part.related_parts[rid]))
    return output


def image_hashes(doc) -> list[str]:
    return [sha256_bytes(part.blob) for _, part in image_parts(doc)]


def document_signature(doc) -> dict:
    return {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "figure_count": len(image_parts(doc)),
        "section_count": len(doc.sections),
    }


def visible_digest(doc) -> str:
    payload = {
        "paragraphs": [paragraph_text(p) for p in doc.paragraphs],
        "tables": [table_matrix(t) for t in doc.tables],
        "images": image_hashes(doc),
        "signature": document_signature(doc),
    }
    return sha256_bytes(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def load_doc(path: Path):
    return Document(path)

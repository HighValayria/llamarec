from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "submission"
TABLES = ROOT / "tables"
FIGURES = ROOT / "figures" / "results"

EN_MD = SUBMISSION / "paper_draft.md"
ZH_MD = SUBMISSION / "paper_draft_zh.md"
TABLES_MD = TABLES / "paper_tables.md"

OUT_EN = SUBMISSION / "llamarec_paper_original_en.docx"
OUT_ZH = SUBMISSION / "llamarec_paper_chinese_zh.docx"


def set_run_font(run, ascii_font: str, east_asia_font: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_font(paragraph, ascii_font: str, east_asia_font: str, size: float | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, ascii_font, east_asia_font, size)


def set_cell_text(cell, text: str, lang: str, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.text = ""
    run = p.add_run(text.strip())
    set_run_font(run, "Calibri", "SimSun" if lang == "zh" else "Calibri", 8.5, bold)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_document(doc: Document, lang: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun" if lang == "zh" else "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei" if lang == "zh" else "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document, title: str, subtitle: str, lang: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_run_font(run, "Calibri", "SimHei" if lang == "zh" else "Calibri", 20, True)
    run.font.color.rgb = RGBColor(32, 36, 42)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(12)
    r2 = p2.add_run(subtitle)
    set_run_font(r2, "Calibri", "SimSun" if lang == "zh" else "Calibri", 10.5, False)
    r2.font.color.rgb = RGBColor(100, 110, 120)


def clean_inline(text: str) -> str:
    return text.replace("`", "").replace("[@", "[@")


def add_paragraph_text(doc: Document, text: str, lang: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(clean_inline(text))
    set_run_font(run, "Calibri", "SimSun" if lang == "zh" else "Calibri", 11)


def split_markdown_blocks(md: str) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []
    for line in md.splitlines():
        if line.strip():
            current.append(line.rstrip())
        else:
            if current:
                blocks.append("\n".join(current))
                current = []
    if current:
        blocks.append("\n".join(current))
    return blocks


def add_markdown(doc: Document, md: str, lang: str, skip_first_title: bool = False) -> None:
    for block in split_markdown_blocks(md):
        if block.startswith("# "):
            text = block[2:].strip()
            if skip_first_title:
                skip_first_title = False
                continue
            doc.add_heading(text, level=1)
            set_paragraph_font(doc.paragraphs[-1], "Calibri", "SimHei" if lang == "zh" else "Calibri")
        elif block.startswith("## "):
            doc.add_heading(block[3:].strip(), level=2)
            set_paragraph_font(doc.paragraphs[-1], "Calibri", "SimHei" if lang == "zh" else "Calibri")
        elif block.startswith("### "):
            doc.add_heading(block[4:].strip(), level=3)
            set_paragraph_font(doc.paragraphs[-1], "Calibri", "SimHei" if lang == "zh" else "Calibri")
        elif block.startswith("- "):
            for item in block.splitlines():
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(clean_inline(item[2:].strip()))
                set_run_font(run, "Calibri", "SimSun" if lang == "zh" else "Calibri", 11)
        else:
            add_paragraph_text(doc, " ".join(line.strip() for line in block.splitlines()), lang)


def parse_md_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    rows = []
    for line in lines:
        parts = [p.strip() for p in line.strip().strip("|").split("|")]
        rows.append(parts)
    header = rows[0]
    body = [r for r in rows[2:] if any(cell for cell in r)]
    return header, body


def add_markdown_table(doc: Document, table_lines: list[str], lang: str) -> None:
    header, body = parse_md_table(table_lines)
    if not header:
        return
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for i, cell in enumerate(table.rows[0].cells):
        set_cell_text(cell, header[i], lang, bold=True)
        shade_cell(cell, "F4F6F9")
        set_cell_margins(cell)
    for row in body:
        cells = table.add_row().cells
        for i, cell in enumerate(cells):
            set_cell_text(cell, row[i] if i < len(row) else "", lang)
            set_cell_margins(cell)
    doc.add_paragraph()


def add_tables_appendix(doc: Document, lang: str) -> None:
    doc.add_page_break()
    doc.add_heading("Paper Tables" if lang == "en" else "论文表格", level=1)
    md = TABLES_MD.read_text(encoding="utf-8")
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("Caption:") or line.startswith("Source CSV:") or line.startswith("Note:"):
            add_paragraph_text(doc, line, lang)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines, lang)
            continue
        i += 1


def add_figures_appendix(doc: Document, lang: str) -> None:
    doc.add_page_break()
    doc.add_heading("Rendered Figures" if lang == "en" else "已渲染图", level=1)
    figures = [
        ("Figure 1. Supervision interfaces for recommendation-tuned LLMs", "图 1. 推荐调优 LLM 的监督接口", FIGURES / "fig1_task_framework.png"),
        ("Figure 2. Exposure-aware comparison between N-K0 and SASRec", "图 2. N-K0 与 SASRec 的曝光量感知比较", FIGURES / "fig2_exposure_aware_sasrec.png"),
        ("Figure 3. Candidate difficulty changes ranking interpretation", "图 3. 候选难度改变排名解释", FIGURES / "fig3_candidate_difficulty.png"),
    ]
    for en_caption, zh_caption, path in figures:
        caption = en_caption if lang == "en" else zh_caption
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(caption)
        set_run_font(run, "Calibri", "SimSun" if lang == "zh" else "Calibri", 10.5, True)
        if path.exists():
            doc.add_picture(str(path), width=Inches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()


def add_footer(doc: Document, lang: str) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "LlamaRec paper package - pre-template draft" if lang == "en" else "LlamaRec 论文包 - 模板前草稿"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(p, "Calibri", "SimSun" if lang == "zh" else "Calibri", 9)


def build_docx(md_path: Path, out_path: Path, lang: str) -> None:
    md = md_path.read_text(encoding="utf-8")
    title = md.splitlines()[0].lstrip("# ").strip()
    subtitle = "Original English Word draft with tables and rendered figures." if lang == "en" else "中文翻译 Word 稿，含表格和已渲染图。"
    doc = Document()
    style_document(doc, lang)
    add_title(doc, title, subtitle, lang)
    add_markdown(doc, md, lang, skip_first_title=True)
    add_tables_appendix(doc, lang)
    add_figures_appendix(doc, lang)
    add_footer(doc, lang)
    core = doc.core_properties
    core.title = title
    core.subject = "LlamaRec paper package"
    core.author = "LlamaRec Paper Writing Stage"
    doc.save(out_path)
    print(f"wrote {out_path}")


def main() -> None:
    build_docx(EN_MD, OUT_EN, "en")
    build_docx(ZH_MD, OUT_ZH, "zh")


if __name__ == "__main__":
    main()
